
"""
STEP 3A-4 — PI WORKFLOW MAP V3 (FINAL EVIDENCE SYNTHESIS)

PURPOSE
-------
Build the Step-3A workflow map from the complete `pi` evidence WITHOUT
pretending that a structural cluster or an OCR phrase is itself a business
process.

The workflow map is derived in this order:

1. Canonical pi population:
      segments.jsonl -> 209 pi segments

2. Observable chronology:
      Step-3A-2 Chronological_Phases
      -> cleaned phase sequences per segment
      -> screenshots/context-only records removed from the workflow order
      -> application transitions retained as structural evidence

3. Full-population recurrence:
      all 209 pi segments
      -> phase presence
      -> pairwise precedence/order
      -> cross-session support

4. Structural pattern evidence:
      Step-3A-2 Workflow_Patterns
      -> used as supporting recurrence evidence only
      -> NOT treated as the business-process taxonomy

5. Semantic evidence:
      Step-3A-3 PI_Semantic_Profiles
      -> joined EXACTLY by (session_id, segment_index)
      -> Japanese + analytical English retained
      -> business object/request/status/action/fields/documents/etc.
      -> never extrapolated beyond directly OCR-covered segments

6. Decision evidence:
      raw decision-event counts + direct OCR decision semantics
      -> reported as decision-oriented evidence
      -> NOT automatically labeled "human judgment"

7. Final workflow interpretation:
      The script generates:
          COMMON_CORE phase evidence
          COMMON_ORDER precedence evidence
          OPTIONAL / VARIANT phases
          semantic evidence by phase/pattern
          representative executions from actual chronological phases

IMPORTANT
---------
This script does NOT:
    - infer a single workflow from one representative execution
    - use screenshot filenames as routes
    - use OCR coverage to estimate subtype prevalence
    - treat application variants as different business processes
    - treat an approval/register/hold button as proof of human judgment
    - overwrite segments.jsonl

OUTPUT
------
ONE Excel workbook only:
    outputs/step3a4_pi_workflow_map_v3.xlsx

SHEETS
------
README
Workflow_Map
Common_Order_Evidence
Segment_Workflow_Sequences
Pattern_Support
Semantic_By_Segment
Semantic_By_Phase
Representative_Executions
Decision_Evidence
Validation

DEFAULT CONSENSUS RULES
-----------------------
A phase is COMMON_CORE if:
    - present in >= 50% of the 209 pi segments, and
    - present in >= 3 distinct sessions.

A phase is FREQUENT if:
    - present in >= 25% of segments.

A pairwise order is COMMON_ORDER if:
    - both phases occur in at least 10 segments, and
    - the first phase precedes the second in >= 70% of those segments.

A pairwise order is FREQUENT_ORDER if:
    - both phases occur in at least 10 segments, and
    - precedence is 55%-69.99%.

These are transparent descriptive thresholds, not optimization scores.
"""

from __future__ import annotations

import re
import math
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import pandas as pd


REPO = Path(__file__).resolve().parents[1]
OUT = REPO / "outputs"
SEGMENTS = REPO / "segments.jsonl"

A1_SEG = OUT / "step3a1_pi_authoritative_segment_evidence_v2.csv"
A1_EVENTS = OUT / "step3a1_pi_authoritative_event_detail_v2.csv"
A1_OCR = OUT / "step3a1_pi_ocr_semantic_evidence_v2.csv"

A2_FILES = sorted(
    OUT.glob("step3a2_pi_chronological_workflow_evidence*.xlsx"),
    key=lambda p: p.stat().st_mtime,
    reverse=True,
)

A3_FILES = sorted(
    OUT.glob("step3a3_pi_ocr_semantic_evidence*.xlsx"),
    key=lambda p: p.stat().st_mtime,
    reverse=True,
)


# ---------------------------------------------------------------------------
# Generic helpers
# ---------------------------------------------------------------------------

def norm(v: Any) -> str:
    if v is None:
        return ""
    try:
        if pd.isna(v):
            return ""
    except Exception:
        pass
    return str(v).strip()


def uniq(values, limit=80):
    out = []
    seen = set()
    for v in values:
        s = norm(v)
        if not s or s in seen:
            continue
        seen.add(s)
        out.append(s)
        if len(out) >= limit:
            break
    return out


def join(values, limit=80):
    return " | ".join(uniq(values, limit))


# ---------------------------------------------------------------------------
# Canonical segment loader
# ---------------------------------------------------------------------------

def load_canonical_pi() -> pd.DataFrame:
    if not SEGMENTS.exists():
        raise FileNotFoundError(f"Missing canonical segments file: {SEGMENTS}")

    rows = []

    import json
    with SEGMENTS.open("r", encoding="utf-8", errors="replace") as f:
        for line_no, line in enumerate(f, 1):
            line = line.strip()
            if not line:
                continue

            obj = json.loads(line)

            if obj.get("label") != "pi":
                continue

            sid = obj.get("session_id")
            start = obj.get("start") or obj.get("start_ts")
            end = obj.get("end") or obj.get("end_ts")

            if not sid or not start or not end:
                raise RuntimeError(
                    f"Canonical pi segment missing key fields at line {line_no}"
                )

            rows.append({
                "session_id": str(sid),
                "start": str(start),
                "end": str(end),
                "canonical_line": line_no,
            })

    # Current canonical segments.jsonl does not contain segment_index.
    # Reconstruct it only for this analysis, chronologically within each session.
    grouped = defaultdict(list)
    for r in rows:
        grouped[r["session_id"]].append(r)

    out = []
    for sid, values in grouped.items():
        values.sort(key=lambda r: (r["start"], r["end"]))
        for idx, r in enumerate(values, 1):
            out.append({
                "session_id": sid,
                "segment_index": idx,
                "segment_start": r["start"],
                "segment_end": r["end"],
                "canonical_line": r["canonical_line"],
            })

    return pd.DataFrame(out).sort_values(
        ["session_id", "segment_index"]
    ).reset_index(drop=True)


# ---------------------------------------------------------------------------
# Source loading
# ---------------------------------------------------------------------------

def load_sources():
    for p in (A1_SEG, A1_EVENTS, A1_OCR):
        if not p.exists():
            raise FileNotFoundError(f"Required 3A-1 file missing: {p}")

    if not A2_FILES:
        raise FileNotFoundError("No Step-3A-2 chronological workbook found.")

    if not A3_FILES:
        raise FileNotFoundError("No Step-3A-3 semantic workbook found.")

    a1_seg = pd.read_csv(A1_SEG)
    a1_events = pd.read_csv(A1_EVENTS)
    a1_ocr = pd.read_csv(A1_OCR)

    a2_path = A2_FILES[0]
    a3_path = A3_FILES[0]

    a2_seg = pd.read_excel(a2_path, sheet_name="PI_Segments")
    a2_phases = pd.read_excel(a2_path, sheet_name="Chronological_Phases")
    a2_patterns = pd.read_excel(a2_path, sheet_name="Workflow_Patterns")

    # Step-3A-2 Representative_Timelines may be empty depending on the version;
    # it is therefore optional. We will reconstruct representatives ourselves.
    try:
        a2_timelines = pd.read_excel(a2_path, sheet_name="Representative_Timelines")
    except Exception:
        a2_timelines = pd.DataFrame()

    a3_joined = pd.read_excel(a3_path, sheet_name="PI_OCR_Joined")
    a3_sem = pd.read_excel(a3_path, sheet_name="PI_Semantic_Profiles")

    for df in (
        a1_seg, a1_events, a1_ocr,
        a2_seg, a2_phases, a2_patterns, a2_timelines,
        a3_joined, a3_sem
    ):
        if "session_id" in df.columns:
            df["session_id"] = df["session_id"].astype(str)
        if "segment_index" in df.columns:
            df["segment_index"] = pd.to_numeric(
                df["segment_index"], errors="coerce"
            ).astype("Int64")

    return (
        a1_seg, a1_events, a1_ocr,
        a2_path, a2_seg, a2_phases, a2_patterns, a2_timelines,
        a3_joined, a3_sem, a3_path,
    )


# ---------------------------------------------------------------------------
# Clean observable phases
# ---------------------------------------------------------------------------

PHASE_MAP = {
    "NAVIGATE_BROWSER": "BROWSER_NAVIGATION",
    "BROWSER_CONTEXT": "BROWSER_WORK",
    "WORK_BROWSER": "BROWSER_WORK",
    "CLICK_BROWSER": "BROWSER_WORK",
    "ENTER_BROWSER_FORM": "FORM_ENTRY",
    "TYPE_BROWSER": "FORM_ENTRY",
    "TEXT_INPUT_COMPLETE": "FORM_ENTRY",
    "BROWSER_FORM": "FORM_ENTRY",

    "CLIPBOARD_TRANSFER": "DATA_TRANSFER",
    "COPY_EDIT_SHORTCUT": "DATA_TRANSFER",
    "TRANSFER": "DATA_TRANSFER",

    "WORK_WORD": "DOCUMENT_WORK",
    "EDIT_WORD": "DOCUMENT_WORK",
    "WORK_EXCEL": "DOCUMENT_WORK",
    "EDIT_EXCEL": "DOCUMENT_WORK",
    "WORK_NOTEPAD": "DOCUMENT_WORK",
    "EDIT_NOTEPAD": "DOCUMENT_WORK",
    "DOCUMENT": "DOCUMENT_WORK",
    "DOCUMENT_CONTEXT": "DOCUMENT_WORK",

    "SWITCH_APPLICATION": "APPLICATION_SWITCH",
    "APP_TRANSITION": "APPLICATION_SWITCH",

    "CLICK_DESKTOP": "DESKTOP_INPUT",
    "TYPE": "DESKTOP_INPUT",
    "SCROLL": "DESKTOP_INPUT",
    "DESKTOP_INPUT": "DESKTOP_INPUT",

    "CONTEXT": "CONTEXT",
    "WINDOW_CHANGE": "CONTEXT",
    "UI_CONTEXT": "CONTEXT",
    "EVIDENCE": "EVIDENCE",
    "SCREENSHOT_CHECKPOINT": "EVIDENCE",
}


def normalize_phase_class(value: Any) -> str:
    base = norm(value).split("[", 1)[0]
    return PHASE_MAP.get(base, base or "UNKNOWN")


def clean_sequence(value: Any) -> list[str]:
    s = norm(value)
    if not s:
        return []

    out = []

    for token in re.split(r"\s*→\s*", s):
        p = normalize_phase_class(token)

        # Evidence checkpoints and pure context transitions should not
        # create their own workflow phases.
        if p in {"EVIDENCE", "CONTEXT", "UNKNOWN"}:
            continue

        # Consecutive identical observable phases are one phase.
        if not out or out[-1] != p:
            out.append(p)

    return out


def build_segment_sequences(a2_seg: pd.DataFrame) -> pd.DataFrame:
    out = a2_seg.copy()

    # Use the phase sequence if available, but clean it independently.
    phase_source = (
        out["chronological_phase_sequence"]
        if "chronological_phase_sequence" in out.columns
        else pd.Series([""] * len(out))
    )

    out["clean_phase_sequence"] = phase_source.map(
        lambda x: " → ".join(clean_sequence(x))
    )
    out["clean_phases_list"] = phase_source.map(clean_sequence)

    # Clean application sequence for a separate tooling signature.
    def clean_apps(value):
        vals = []
        for token in re.split(r"\s*→\s*", norm(value)):
            low = token.lower()
            if any(x in low for x in ("edge", "chrome", "firefox", "browser")):
                a = "Browser"
            elif "word" in low:
                a = "Word"
            elif "excel" in low:
                a = "Excel"
            elif "notepad" in low:
                a = "Notepad"
            elif "explorer" in low:
                a = "Explorer"
            elif "openwith" in low:
                a = "OpenWith"
            elif "terminal" in low or "powershell" in low or "cmd" in low:
                a = "Terminal"
            else:
                a = "OtherApp"

            if not vals or vals[-1] != a:
                vals.append(a)

        return " → ".join(vals)

    app_col = (
        "application_sequence_raw_events"
        if "application_sequence_raw_events" in out.columns
        else "application_sequence"
    )

    out["clean_application_sequence"] = out[app_col].map(clean_apps)

    out["has_clipboard"] = (
        pd.to_numeric(
            out.get("clipboard_events", pd.Series([0] * len(out))),
            errors="coerce",
        ).fillna(0).gt(0)
    )

    out["has_form_entry"] = (
        pd.to_numeric(
            out.get("browser_form_input_events", pd.Series([0] * len(out))),
            errors="coerce",
        ).fillna(0).gt(0)
    )

    out["has_navigation"] = (
        pd.to_numeric(
            out.get("browser_navigation_events", pd.Series([0] * len(out))),
            errors="coerce",
        ).fillna(0).gt(0)
    )

    out["has_decision_evidence"] = (
        pd.to_numeric(
            out.get("decision_point_raw_event_count", pd.Series([0] * len(out))),
            errors="coerce",
        ).fillna(0).gt(0)
    )

    out["has_document"] = (
        out["clean_application_sequence"]
        .fillna("")
        .str.contains(r"Word|Excel|Notepad|Explorer|OpenWith", regex=True)
    )

    return out


# ---------------------------------------------------------------------------
# Pattern support
# ---------------------------------------------------------------------------

def build_pattern_support(
    a2_patterns: pd.DataFrame,
    segment_df: pd.DataFrame,
) -> pd.DataFrame:

    p = a2_patterns.copy()

    # The A2 pattern table may include singleton patterns. Keep all for
    # transparency, but identify the recurring/cross-session subset.
    p["pattern_evidence_class"] = "SINGLETON"
    p.loc[
        (p["segment_count"].astype(int) >= 2)
        & (p["session_count"].astype(int) >= 2),
        "pattern_evidence_class"
    ] = "CROSS_SESSION_RECURRING"
    p.loc[
        (p["segment_count"].astype(int) >= 2)
        & (p["session_count"].astype(int) < 2),
        "pattern_evidence_class"
    ] = "WITHIN_SESSION_RECURRING"

    return p.sort_values(
        ["segment_count", "session_count"],
        ascending=[False, False]
    ).reset_index(drop=True)


# ---------------------------------------------------------------------------
# Phase support across all 209 segments
# ---------------------------------------------------------------------------

def build_phase_support(segment_df: pd.DataFrame) -> pd.DataFrame:
    rows = []

    total_segments = len(segment_df)

    # Map phase -> segment/session sets.
    segs = defaultdict(set)
    sessions = defaultdict(set)
    positions = defaultdict(list)

    for _, r in segment_df.iterrows():
        sid = str(r["session_id"])
        idx = int(r["segment_index"])
        seq = r["clean_phases_list"]

        for pos, phase in enumerate(seq, 1):
            key = phase
            segs[key].add((sid, idx))
            sessions[key].add(sid)
            positions[key].append(pos)

    for phase, keys in segs.items():
        count = len(keys)
        sess = len(sessions[phase])
        pct = 100 * count / total_segments if total_segments else 0

        if pct >= 50 and sess >= 3:
            classification = "COMMON_CORE"
        elif pct >= 25:
            classification = "FREQUENT"
        else:
            classification = "OPTIONAL_VARIANT"

        rows.append({
            "phase": phase,
            "segment_count": count,
            "segment_coverage_pct": round(pct, 2),
            "session_count": sess,
            "median_position": float(pd.Series(positions[phase]).median()),
            "mean_position": float(pd.Series(positions[phase]).mean()),
            "classification": classification,
        })

    return pd.DataFrame(rows).sort_values(
        ["classification", "segment_count"],
        ascending=[True, False]
    ).reset_index(drop=True)


# ---------------------------------------------------------------------------
# Pairwise precedence across all segments
# ---------------------------------------------------------------------------

def build_order_evidence(segment_df: pd.DataFrame) -> pd.DataFrame:
    pair_counts = Counter()
    first_counts = Counter()
    phase_presence = Counter()

    for _, r in segment_df.iterrows():
        seq = r["clean_phases_list"]

        # Only once per segment.
        unique_seq = []
        for x in seq:
            if x not in unique_seq:
                unique_seq.append(x)

        for p in unique_seq:
            phase_presence[p] += 1

        for i, a in enumerate(unique_seq):
            for b in unique_seq[i + 1:]:
                pair_counts[(a, b)] += 1
                first_counts[(a, b)] += 1

    # We also need reverse-pair counts to calculate precedence.
    co_presence = Counter()

    all_phases = sorted(phase_presence)
    for i, a in enumerate(all_phases):
        for b in all_phases[i + 1:]:
            co = 0
            a_before_b = 0
            b_before_a = 0

            for _, r in segment_df.iterrows():
                seq = []
                for x in r["clean_phases_list"]:
                    if x not in seq:
                        seq.append(x)

                if a in seq and b in seq:
                    co += 1
                    if seq.index(a) < seq.index(b):
                        a_before_b += 1
                    elif seq.index(b) < seq.index(a):
                        b_before_a += 1

            if co == 0:
                continue

            # Canonicalize orientation to the direction with greater support.
            if a_before_b >= b_before_a:
                first, second = a, b
                precedence = a_before_b / co
            else:
                first, second = b, a
                precedence = b_before_a / co

            if co >= 10 and precedence >= 0.70:
                cls = "COMMON_ORDER"
            elif co >= 10 and precedence >= 0.55:
                cls = "FREQUENT_ORDER"
            else:
                cls = "MIXED_ORDER"

            order_rows = {
                "first_phase": first,
                "second_phase": second,
                "segments_with_both": co,
                "first_before_second_count": max(a_before_b, b_before_a),
                "second_before_first_count": min(a_before_b, b_before_a),
                "precedence_pct": round(100 * precedence, 2),
                "classification": cls,
            }
            co_presence[(first, second)] = order_rows

    return pd.DataFrame(
        list(co_presence.values())
    ).sort_values(
        ["classification", "segments_with_both", "precedence_pct"],
        ascending=[True, False, False]
    ).reset_index(drop=True)


# ---------------------------------------------------------------------------
# Derive consensus structural order
# ---------------------------------------------------------------------------

def build_consensus_order(
    phase_support: pd.DataFrame,
    order_evidence: pd.DataFrame,
) -> pd.DataFrame:

    common_phases = phase_support[
        phase_support["classification"] == "COMMON_CORE"
    ]["phase"].tolist()

    if not common_phases:
        return pd.DataFrame(columns=[
            "consensus_rank",
            "phase",
            "support_pct",
            "session_count",
            "order_score",
            "evidence_note",
        ])

    # Calculate a precedence score:
    # sum over strong/common relationships of +1 if phase precedes another,
    # -1 if it follows. This is descriptive, not a ranking score used to choose
    # a process candidate.
    order_score = Counter()

    for phase in common_phases:
        order_score[phase] = 0

    for _, r in order_evidence.iterrows():
        if r["classification"] != "COMMON_ORDER":
            continue
        a = norm(r["first_phase"])
        b = norm(r["second_phase"])
        if a in common_phases and b in common_phases:
            order_score[a] += 1
            order_score[b] -= 1

    rows = []

    for _, r in phase_support.iterrows():
        if r["phase"] not in common_phases:
            continue

        rows.append({
            "phase": r["phase"],
            "support_pct": r["segment_coverage_pct"],
            "session_count": r["session_count"],
            "order_score": order_score[r["phase"]],
            "median_position": r["median_position"],
            "evidence_note":
                "Consensus position derived from common phases and common pairwise precedence; not a directly observed single universal sequence.",
        })

    out = pd.DataFrame(rows).sort_values(
        ["order_score", "median_position", "support_pct"],
        ascending=[False, True, False]
    ).reset_index(drop=True)

    out.insert(0, "consensus_rank", range(1, len(out) + 1))
    return out


# ---------------------------------------------------------------------------
# Exact OCR semantic join
# ---------------------------------------------------------------------------

def build_semantic_by_segment(
    segment_df: pd.DataFrame,
    a3_sem: pd.DataFrame,
    a3_joined: pd.DataFrame,
) -> pd.DataFrame:
    """
    Exact segment-level semantic join.

    Semantic fields come from PI_Semantic_Profiles.
    OCR Japanese/English evidence comes from PI_OCR_Joined.
    Both are joined by (session_id, segment_index).
    """
    base = segment_df[
        [
            "session_id",
            "segment_index",
            "clean_phase_sequence",
            "clean_application_sequence",
            "has_clipboard",
            "has_form_entry",
            "has_decision_evidence",
        ]
    ].copy()

    # ---- Semantic profile table ----
    sem = a3_sem.copy()
    if not sem.empty:
        sem["session_id"] = sem["session_id"].astype(str)
        sem["segment_index"] = pd.to_numeric(
            sem["segment_index"], errors="coerce"
        ).astype("Int64")

        semantic_columns = [
            "workflow_signature",
            "business_object",
            "request_type",
            "field_names",
            "status",
            "action",
            "document",
            "visible_instruction",
            "decision_options",
        ]
        available = [c for c in semantic_columns if c in sem.columns]

        # Collapse possible duplicate OCR-derived semantic rows per segment.
        agg_map = {
            c: (lambda x, _c=c: join(x.tolist(), 100))
            for c in available
        }

        sem_grouped = (
            sem.groupby(
                ["session_id", "segment_index"],
                dropna=False
            )
            .agg(agg_map)
            .reset_index()
        )
    else:
        sem_grouped = pd.DataFrame(
            columns=["session_id", "segment_index"]
        )

    # ---- OCR text evidence table ----
    ocr = a3_joined.copy()
    if not ocr.empty:
        ocr["session_id"] = ocr["session_id"].astype(str)
        ocr["segment_index"] = pd.to_numeric(
            ocr["segment_index"], errors="coerce"
        ).astype("Int64")

        ocr_cols = [
            c for c in (
                "ocr_japanese",
                "ocr_english",
                "ocr_record_count",
                "ocr_mean_score",
            )
            if c in ocr.columns
        ]

        ocr_agg = {}

        if "ocr_japanese" in ocr.columns:
            ocr_agg["ocr_japanese"] = lambda x: join(x.tolist(), 120)

        if "ocr_english" in ocr.columns:
            ocr_agg["ocr_english"] = lambda x: join(x.tolist(), 120)

        if "ocr_record_count" in ocr.columns:
            ocr_agg["ocr_record_count"] = lambda x: (
                pd.to_numeric(x, errors="coerce").fillna(0).max()
            )

        if "ocr_mean_score" in ocr.columns:
            ocr_agg["ocr_mean_score"] = lambda x: (
                pd.to_numeric(x, errors="coerce").dropna().mean()
                if pd.to_numeric(x, errors="coerce").notna().any()
                else ""
            )

        if ocr_agg:
            ocr_grouped = (
                ocr.groupby(
                    ["session_id", "segment_index"],
                    dropna=False
                )
                .agg(ocr_agg)
                .reset_index()
            )
        else:
            ocr_grouped = pd.DataFrame(
                columns=["session_id", "segment_index"]
            )
    else:
        ocr_grouped = pd.DataFrame(
            columns=["session_id", "segment_index"]
        )

    out = base.merge(
        sem_grouped,
        on=["session_id", "segment_index"],
        how="left",
        suffixes=("", "_semantic"),
    )

    out = out.merge(
        ocr_grouped,
        on=["session_id", "segment_index"],
        how="left",
    )

    for col in (
        "workflow_signature",
        "business_object",
        "request_type",
        "field_names",
        "status",
        "action",
        "document",
        "visible_instruction",
        "decision_options",
        "ocr_japanese",
        "ocr_english",
    ):
        if col not in out.columns:
            out[col] = ""
        out[col] = out[col].fillna("").astype(str)

    if "ocr_record_count" not in out.columns:
        out["ocr_record_count"] = 0

    out["direct_ocr_present"] = (
        out["ocr_record_count"]
        .apply(lambda x: (float(x) if str(x).strip() else 0.0) > 0)
    )

    return out


# ---------------------------------------------------------------------------
# Semantic summaries by phase
# ---------------------------------------------------------------------------

def build_semantic_by_phase(
    semantic_by_segment: pd.DataFrame,
) -> pd.DataFrame:

    # Direct OCR only.
    direct = semantic_by_segment[
        semantic_by_segment["direct_ocr_present"]
    ].copy()

    if direct.empty:
        return pd.DataFrame()

    rows = []

    for phase in sorted(
        set(
            x.strip()
            for s in direct["clean_phase_sequence"]
            for x in norm(s).split("→")
            if x.strip()
        )
    ):
        g = direct[
            direct["clean_phase_sequence"]
            .astype(str)
            .str.contains(
                rf"(?:^|→)\s*{re.escape(phase)}(?:\s*→|$)",
                regex=True,
            )
        ]

        if g.empty:
            continue

        rows.append({
            "phase": phase,
            "direct_ocr_segments": len(g),
            "business_objects_observed": join(g["business_object"]),
            "request_types_observed": join(g["request_type"]),
            "fields_observed": join(g["field_names"]),
            "statuses_observed": join(g["status"]),
            "actions_observed": join(g["action"]),
            "documents_observed": join(g["document"]),
            "instructions_observed": join(g["visible_instruction"]),
            "decision_options_observed": join(g["decision_options"]),
            "japanese_evidence": join(g["ocr_japanese"], 120),
            "english_evidence": join(g["ocr_english"], 120),
            "interpretation_status":
                "Direct OCR semantic evidence only; not extrapolated to uncovered segments.",
        })

    return pd.DataFrame(rows)


# ---------------------------------------------------------------------------
# Representative executions
# ---------------------------------------------------------------------------

def choose_representatives(
    segment_df: pd.DataFrame,
    semantic_by_segment: pd.DataFrame,
    order_evidence: pd.DataFrame,
    n=12,
) -> pd.DataFrame:

    candidates = segment_df.copy()

    candidates["ocr_bonus"] = (
        semantic_by_segment["direct_ocr_present"].astype(int)
    ).values

    # Prefer:
    #   - direct OCR
    #   - decision evidence
    #   - cross-application behavior
    #   - a duration close to the median
    # while avoiding selecting many executions from one session.
    med_duration = pd.to_numeric(
        candidates["duration_seconds"],
        errors="coerce"
    ).median()

    candidates["decision_bonus"] = candidates[
        "has_decision_evidence"
    ].astype(int)

    candidates["cross_app_bonus"] = (
        candidates["clean_application_sequence"]
        .fillna("")
        .str.contains("→", regex=False)
        .astype(int)
    )

    candidates["representative_distance"] = (
        pd.to_numeric(
            candidates["duration_seconds"],
            errors="coerce"
        ).fillna(med_duration) - med_duration
    ).abs()

    candidates["rep_score"] = (
        candidates["ocr_bonus"] * 1000
        + candidates["decision_bonus"] * 100
        + candidates["cross_app_bonus"] * 10
        - candidates["representative_distance"]
    )

    selected = []
    used_sessions = set()

    for _, row in candidates.sort_values(
        "rep_score", ascending=False
    ).iterrows():
        sid = str(row["session_id"])

        if sid in used_sessions:
            continue

        selected.append(row)
        used_sessions.add(sid)

        if len(selected) >= n:
            break

    # Fill remaining slots if fewer than n sessions were available.
    if len(selected) < min(n, len(candidates)):
        selected_keys = {
            (str(x["session_id"]), int(x["segment_index"]))
            for x in selected
        }

        for _, row in candidates.sort_values(
            "rep_score", ascending=False
        ).iterrows():
            key = (str(row["session_id"]), int(row["segment_index"]))
            if key in selected_keys:
                continue

            selected.append(row)
            selected_keys.add(key)

            if len(selected) >= n:
                break

    return pd.DataFrame(selected)[[
        "session_id",
        "segment_index",
        "segment_start",
        "segment_end",
        "duration_seconds",
        "raw_event_count",
        "clean_application_sequence",
        "clean_phase_sequence",
        "has_clipboard",
        "has_form_entry",
        "has_decision_evidence",
        "has_document",
    ]].reset_index(drop=True)


# ---------------------------------------------------------------------------
# Decision evidence
# ---------------------------------------------------------------------------

def build_decision_evidence(
    a1_seg: pd.DataFrame,
    a3_sem: pd.DataFrame,
) -> pd.DataFrame:

    raw = a1_seg.copy()

    if "decision_point_raw_event_count" in raw.columns:
        raw["decision_count"] = pd.to_numeric(
            raw["decision_point_raw_event_count"],
            errors="coerce"
        ).fillna(0).astype(int)
    else:
        raw["decision_count"] = 0

    raw = raw[raw["decision_count"] > 0].copy()

    rows = []

    for _, r in raw.iterrows():
        rows.append({
            "session_id": str(r["session_id"]),
            "segment_index": int(r["segment_index"]),
            "evidence_source": "RAW_EVENT",
            "decision_evidence_count": int(r["decision_count"]),
            "observed_signal": norm(r.get("decision_event_types")),
            "japanese_evidence": "",
            "english_evidence": "",
            "interpretation":
                "Decision-oriented event evidence observed; human judgment requirement remains unresolved without contextual review.",
        })

    if not a3_sem.empty:
        for _, r in a3_sem.iterrows():
            actions = norm(r.get("action"))
            options = norm(r.get("decision_options"))

            if actions or options:
                rows.append({
                    "session_id": str(r["session_id"]),
                    "segment_index": int(r["segment_index"]),
                    "evidence_source": "DIRECT_OCR",
                    "decision_evidence_count": "",
                    "observed_signal":
                        "; ".join(x for x in (actions, options) if x),
                    "japanese_evidence": norm(r.get("ocr_japanese")),
                    "english_evidence": norm(r.get("ocr_english")),
                    "interpretation":
                        "Decision-oriented UI/action semantics observed directly in OCR; button/action text alone does not establish required human judgment.",
                })

    if not rows:
        return pd.DataFrame(columns=[
            "session_id", "segment_index", "evidence_source",
            "decision_evidence_count", "observed_signal",
            "japanese_evidence", "english_evidence", "interpretation"
        ])

    return pd.DataFrame(rows).drop_duplicates()


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------

def build_validation(
    canonical: pd.DataFrame,
    a2_seg: pd.DataFrame,
    a2_phases: pd.DataFrame,
    a2_patterns: pd.DataFrame,
    phase_support: pd.DataFrame,
    order_evidence: pd.DataFrame,
    semantic: pd.DataFrame,
) -> pd.DataFrame:

    canonical_keys = set(
        (str(s), int(i))
        for s, i in zip(
            canonical["session_id"], canonical["segment_index"]
        )
    )

    a2_keys = set(
        (str(s), int(i))
        for s, i in zip(
            a2_seg["session_id"], a2_seg["segment_index"]
        )
    )

    rows = [
        {
            "check": "canonical_pi_segment_count",
            "expected": 209,
            "observed": len(canonical),
            "status": "PASS" if len(canonical) == 209 else "FAIL",
        },
        {
            "check": "A2_pi_segment_count",
            "expected": 209,
            "observed": len(a2_seg),
            "status": "PASS" if len(a2_seg) == 209 else "FAIL",
        },
        {
            "check": "canonical_A2_segment_key_match",
            "expected": 209,
            "observed": len(canonical_keys & a2_keys),
            "status": (
                "PASS"
                if len(canonical_keys & a2_keys) == 209
                else "FAIL"
            ),
        },
        {
            "check": "A2_chronological_phase_rows",
            "expected": ">0",
            "observed": len(a2_phases),
            "status": "PASS" if len(a2_phases) > 0 else "FAIL",
        },
        {
            "check": "A2_workflow_patterns",
            "expected": ">0",
            "observed": len(a2_patterns),
            "status": "PASS" if len(a2_patterns) > 0 else "FAIL",
        },
        {
            "check": "common_core_phases",
            "expected": ">0",
            "observed": int(
                (phase_support["classification"] == "COMMON_CORE").sum()
            ) if not phase_support.empty else 0,
            "status": (
                "PASS"
                if not phase_support.empty
                and int(
                    (phase_support["classification"] == "COMMON_CORE").sum()
                ) > 0
                else "REVIEW"
            ),
        },
        {
            "check": "order_evidence_rows",
            "expected": ">0",
            "observed": len(order_evidence),
            "status": "PASS" if len(order_evidence) > 0 else "REVIEW",
        },
        {
            "check": "direct_ocr_semantic_segments",
            "expected": "",
            "observed": int(semantic["direct_ocr_present"].sum())
            if not semantic.empty else 0,
            "status": "INFO",
        },
    ]

    return pd.DataFrame(rows)


# ---------------------------------------------------------------------------
# Excel output
# ---------------------------------------------------------------------------

def write_excel(
    path: Path,
    readme: pd.DataFrame,
    workflow_map: pd.DataFrame,
    order_evidence: pd.DataFrame,
    segment_sequences: pd.DataFrame,
    pattern_support: pd.DataFrame,
    semantic_by_segment: pd.DataFrame,
    semantic_by_phase: pd.DataFrame,
    representatives: pd.DataFrame,
    decision_evidence: pd.DataFrame,
    validation: pd.DataFrame,
):

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        readme.to_excel(writer, sheet_name="README", index=False)
        workflow_map.to_excel(writer, sheet_name="Workflow_Map", index=False)
        order_evidence.to_excel(
            writer, sheet_name="Common_Order_Evidence", index=False
        )
        segment_sequences.to_excel(
            writer, sheet_name="Segment_Workflow_Sequences", index=False
        )
        pattern_support.to_excel(
            writer, sheet_name="Pattern_Support", index=False
        )
        semantic_by_segment.to_excel(
            writer, sheet_name="Semantic_By_Segment", index=False
        )
        semantic_by_phase.to_excel(
            writer, sheet_name="Semantic_By_Phase", index=False
        )
        representatives.to_excel(
            writer, sheet_name="Representative_Executions", index=False
        )
        decision_evidence.to_excel(
            writer, sheet_name="Decision_Evidence", index=False
        )
        validation.to_excel(writer, sheet_name="Validation", index=False)

        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        wb = writer.book

        for ws in wb.worksheets:
            ws.freeze_panes = "A2"
            if ws.max_row > 1 and ws.max_column > 1:
                ws.auto_filter.ref = ws.dimensions

            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(
                    "solid",
                    fgColor="1F4E78"
                )
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )

            for col_idx, cells in enumerate(
                ws.iter_cols(
                    min_row=1,
                    max_row=min(ws.max_row, 120)
                ),
                start=1,
            ):
                max_len = max(
                    [len(str(c.value or "")) for c in cells] + [12]
                )
                ws.column_dimensions[
                    get_column_letter(col_idx)
                ].width = min(max_len + 2, 48)

            for row in range(2, min(ws.max_row, 300) + 1):
                ws.row_dimensions[row].height = 30


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():

    print("=" * 80)
    print("STEP 3A-4 — PI WORKFLOW MAP V3 (FINAL EVIDENCE SYNTHESIS)")
    print("=" * 80)

    (
        a1_seg,
        a1_events,
        a1_ocr,
        a2_path,
        a2_seg,
        a2_phases,
        a2_patterns,
        a2_timelines,
        a3_joined,
        a3_sem,
        a3_path,
    ) = load_sources()

    canonical = load_canonical_pi()

    print(f"Canonical pi segments: {len(canonical)}")
    print(f"3A-2 segments: {len(a2_seg)}")
    print(f"3A-2 phases: {len(a2_phases)}")
    print(f"3A-2 patterns: {len(a2_patterns)}")
    print(f"3A-3 semantic profiles: {len(a3_sem)}")
    print(f"Using A2 workbook: {a2_path.name}")
    print(f"Using A3 workbook: {a3_path.name}")

    # Ensure exact 209-segment foundation.
    if len(canonical) != 209:
        raise RuntimeError(
            f"Canonical pi population is {len(canonical)}, expected 209."
        )

    segment_df = build_segment_sequences(a2_seg)

    if len(segment_df) != 209:
        raise RuntimeError(
            f"Step-3A-2 pi segment population is {len(segment_df)}, expected 209."
        )

    # Phase support and pairwise order are derived from ALL 209 segments.
    phase_support = build_phase_support(segment_df)
    order_evidence = build_order_evidence(segment_df)
    consensus = build_consensus_order(
        phase_support,
        order_evidence,
    )

    # Structural pattern support is descriptive only.
    pattern_support = build_pattern_support(
        a2_patterns,
        segment_df
    )

    # Exact OCR semantic join.
    semantic_by_segment = build_semantic_by_segment(
        segment_df,
        a3_sem,
        a3_joined,
    )

    semantic_by_phase = build_semantic_by_phase(
        semantic_by_segment
    )

    # Representative executions are selected from actual segment chronology,
    # not from a synthetic "representative timeline" sheet.
    representatives = choose_representatives(
        segment_df,
        semantic_by_segment,
        order_evidence,
        n=12,
    )

    decision = build_decision_evidence(
        a1_seg,
        a3_sem,
    )

    # Final workflow map = consensus structural order + common support + semantic
    # meaning where directly evidenced.
    workflow_map = consensus.merge(
        phase_support[[
            "phase",
            "segment_count",
            "segment_coverage_pct",
            "session_count",
            "classification",
        ]],
        on="phase",
        how="left",
        suffixes=("", "_support"),
    )

    semantic_phase_lookup = (
        semantic_by_phase.set_index("phase")
        if not semantic_by_phase.empty and "phase" in semantic_by_phase.columns
        else pd.DataFrame()
    )

    if not semantic_phase_lookup.empty:
        for col in (
            "business_objects_observed",
            "request_types_observed",
            "fields_observed",
            "statuses_observed",
            "actions_observed",
            "documents_observed",
            "instructions_observed",
            "decision_options_observed",
        ):
            workflow_map[col] = workflow_map["phase"].map(
                semantic_phase_lookup[col]
            ).fillna("")
    else:
        for col in (
            "business_objects_observed",
            "request_types_observed",
            "fields_observed",
            "statuses_observed",
            "actions_observed",
            "documents_observed",
            "instructions_observed",
            "decision_options_observed",
        ):
            workflow_map[col] = ""

    # Add a clear interpretation category.
    workflow_map["interpretation"] = workflow_map["phase"].map({
        "BROWSER_NAVIGATION":
            "Navigation into/within the browser application.",
        "BROWSER_WORK":
            "Observable work in the browser/target UI.",
        "DATA_TRANSFER":
            "Observable data transfer/copy activity.",
        "DOCUMENT_WORK":
            "Observable document preparation/editing.",
        "APPLICATION_SWITCH":
            "Movement between applications/contexts.",
        "FORM_ENTRY":
            "Observable structured browser form/field entry.",
        "DESKTOP_INPUT":
            "Observable keyboard/mouse/scroll activity not otherwise assigned.",
    }).fillna(
        "Observable structural activity; business meaning requires contextual evidence."
    )

    validation = build_validation(
        canonical,
        a2_seg,
        a2_phases,
        a2_patterns,
        phase_support,
        order_evidence,
        semantic_by_segment,
    )

    readme = pd.DataFrame([
        [
            "Purpose",
            "Final evidence synthesis for reconstructing the pi workflow before deciding the automation boundary.",
        ],
        [
            "Population",
            "All 209 canonical pi segments are included in the structural analysis.",
        ],
        [
            "Chronology",
            "Phase order is derived from Step-3A-2 chronological phases after removing screenshot/context-only noise.",
        ],
        [
            "Common phase rule",
            ">=50% of pi segments and >=3 sessions.",
        ],
        [
            "Common order rule",
            "Both phases in >=10 segments and precedence >=70%.",
        ],
        [
            "Semantic evidence",
            "Direct OCR semantics are joined by exact (session_id, segment_index). No semantic prevalence is extrapolated.",
        ],
        [
            "Pattern role",
            "Step-3A-2 structural patterns are supporting evidence, not final business-process labels.",
        ],
        [
            "Decision evidence",
            "Decision-oriented events/UI semantics are reported separately; labels such as approve/register/hold do not by themselves prove human judgment.",
        ],
        [
            "Output status",
            "Evidence-backed workflow map. This is still pre-automation-design.",
        ],
        [
            "Segmentation",
            "segments.jsonl is never modified.",
        ],
    ], columns=["item", "value"])

    output = OUT / "step3a4_pi_workflow_map_v3.xlsx"

    write_excel(
        output,
        readme,
        workflow_map,
        order_evidence,
        segment_df[[
            "session_id",
            "segment_index",
            "segment_start",
            "segment_end",
            "duration_seconds",
            "raw_event_count",
            "clean_application_sequence",
            "clean_phase_sequence",
            "has_clipboard",
            "has_form_entry",
            "has_navigation",
            "has_document",
            "has_decision_evidence",
        ]],
        pattern_support,
        semantic_by_segment,
        semantic_by_phase,
        representatives,
        decision,
        validation,
    )

    print("\n" + "=" * 80)
    print("3A-4 RESULT")
    print("=" * 80)
    print(f"Canonical pi segments: {len(canonical)}")
    print(
        f"COMMON_CORE phases: "
        f"{int((phase_support['classification'] == 'COMMON_CORE').sum())}"
        if not phase_support.empty else
        "COMMON_CORE phases: 0"
    )
    print(
        f"Common-order relationships: "
        f"{int((order_evidence['classification'] == 'COMMON_ORDER').sum())}"
        if not order_evidence.empty else
        "Common-order relationships: 0"
    )
    print(
        f"Direct OCR semantic segments: "
        f"{int(semantic_by_segment['direct_ocr_present'].sum())}"
        if not semantic_by_segment.empty else
        "Direct OCR semantic segments: 0"
    )
    print(f"Representative executions: {len(representatives)}")
    print(f"Decision evidence rows: {len(decision)}")

    print("\nCONSENSUS WORKFLOW ORDER")
    if consensus.empty:
        print("No defensible common-core order was established.")
    else:
        print(
            consensus.to_string(index=False)
        )

    print("\nCOMMON ORDER EVIDENCE")
    if order_evidence.empty:
        print("No pairwise order evidence.")
    else:
        print(
            order_evidence[
                order_evidence["classification"].isin(
                    ["COMMON_ORDER", "FREQUENT_ORDER"]
                )
            ].head(25).to_string(index=False)
        )

    print("\nVALIDATION")
    print(validation.to_string(index=False))

    failures = validation[validation["status"] == "FAIL"]
    if not failures.empty:
        raise RuntimeError(
            "3A-4 validation failed. Inspect the Validation sheet."
        )

    print("\nOUTPUT")
    print(output)
    print("\nsegments.jsonl was NOT modified.")
    print("STEP 3A-4 V3 COMPLETE.")


if __name__ == "__main__":
    main()
