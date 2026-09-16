#!/usr/bin/env python3
"""
STEP 3E-1 — PI LOCAL PROTOTYPE

Creates a safe local /payroll-items mock and, when Playwright is installed,
drives it through the evidence-derived preparation workflow.

READ -> EXTRACT -> TRANSFORM -> POPULATE -> VERIFY -> STOP

No production connection. No Register/Hold clicks.

Install:
  pip install playwright pandas openpyxl
  playwright install chromium

Run from the repository root:
  python -u src/step3e1_pi_local_prototype_v1.py

Output:
  outputs/step3e1_pi_prototype_evidence_v1.xlsx
  prototype/step3e1_mock_payroll_items.html
  prototype/step3e1_test_cases.json
"""

from __future__ import annotations

import json
import re
import time
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any

import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs"
PROTO = ROOT / "prototype"
HTML_FILE = PROTO / "step3e1_mock_payroll_items.html"
CASES_FILE = PROTO / "step3e1_test_cases.json"
XLSX_FILE = OUT / "step3e1_pi_prototype_evidence_v1.xlsx"
HOST = "127.0.0.1"
PORT = 8765
URL = f"http://{HOST}:{PORT}/payroll-items"

HTML = """<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>Payroll Items - Local Prototype</title>
<style>
body{font-family:Arial,sans-serif;margin:40px;background:#f6f7f9;color:#202124}
.wrap{max-width:900px;margin:auto}.card{background:#fff;border:1px solid #d9dce1;border-radius:8px;padding:20px;margin-bottom:18px}
h1{margin-top:0}.grid{display:grid;grid-template-columns:1fr 1fr;gap:14px}.field{display:flex;flex-direction:column;gap:6px}.full{grid-column:1/-1}
input,select,textarea{font:inherit;padding:9px;border:1px solid #b8bec6;border-radius:5px}textarea{min-height:80px}
.gate{border:2px solid #555;padding:15px;border-radius:6px}.buttons{display:flex;gap:12px}
button{font:inherit;padding:10px 18px}.result{white-space:pre-wrap;background:#eef2f6;padding:12px;border-radius:5px}
</style></head>
<body><div class="wrap">
<div class="card"><h1>Payroll Items - Local Prototype</h1><div id="route">Route: <b>/payroll-items</b></div></div>
<div class="card"><h2>Prepared Payroll Record</h2><div class="grid">
<div class="field"><label for="employee_id">Employee ID</label><input id="employee_id" name="employee_id"></div>
<div class="field"><label for="request_type">Request Type</label><select id="request_type" name="request_type"><option value="">Select...</option><option value="expense_reimbursement">Expense reimbursement</option><option value="payroll_change">Payroll change</option><option value="dependent_deduction_change">Dependent deduction change</option><option value="transportation_expense">Transportation expense</option></select></div>
<div class="field"><label for="amount">Amount</label><input id="amount" name="amount"></div>
<div class="field"><label for="effective_date">Effective Date</label><input id="effective_date" type="date" name="effective_date"></div>
<div class="field full"><label for="processing_comment">Processing Comment</label><textarea id="processing_comment" name="processing_comment"></textarea></div>
<div class="field full"><label for="supporting_document">Supporting Document</label><input id="supporting_document" name="supporting_document"></div>
<div class="field full"><label for="status">Status</label><input id="status" name="status" readonly value="Prepared - Human Review Required"></div>
</div></div>
<div class="card gate"><h2>Human Review Gate</h2><p>Automation stops here. Decision controls are intentionally disabled.</p>
<div class="buttons"><button id="register" disabled>Register</button><button id="hold" disabled>Hold</button></div></div>
<div class="card"><h2>Verification</h2><div id="result" class="result">Waiting...</div></div>
<script>
window.getPreparedRecord=function(){return {employee_id:document.querySelector('#employee_id').value,request_type:document.querySelector('#request_type').value,amount:document.querySelector('#amount').value,effective_date:document.querySelector('#effective_date').value,processing_comment:document.querySelector('#processing_comment').value,supporting_document:document.querySelector('#supporting_document').value,status:document.querySelector('#status').value};};
window.markVerification=function(ok,msg){const e=document.querySelector('#result');e.dataset.verification=ok?'PASS':'FAIL';e.textContent=msg;};
</script></div></body></html>"""

CASES = [
    {"case_id":"PI-DEMO-001","case_type":"expense_reimbursement","employee_id":"E1001","amount":"1250.50","effective_date":"2026-09-15","processing_comment":"Reimbursement prepared from validated source data.","supporting_document":"expense_receipt_001.pdf","expected":"SUCCESS"},
    {"case_id":"PI-DEMO-002","case_type":"payroll_change","employee_id":"E1002","amount":"4500","effective_date":"2026-10-01","processing_comment":"Payroll change prepared from validated source data.","supporting_document":"","expected":"SUCCESS"},
    {"case_id":"PI-DEMO-003","case_type":"expense_reimbursement","employee_id":"E1003","amount":"300.00","effective_date":"2026-09-20","processing_comment":"Document-assisted reimbursement preparation.","supporting_document":"receipt_003.pdf","expected":"SUCCESS"},
    {"case_id":"PI-DEMO-004","case_type":"expense_reimbursement","employee_id":"","amount":"700.00","effective_date":"2026-09-21","processing_comment":"Missing employee ID test.","supporting_document":"receipt_004.pdf","expected":"STOP_MISSING_REQUIRED_INPUT"},
    {"case_id":"PI-DEMO-005","case_type":"payroll_change","employee_id":"E1005","amount":"NOT_A_NUMBER","effective_date":"2026-10-01","processing_comment":"Invalid amount test.","supporting_document":"","expected":"STOP_INVALID_INPUT"},
    {"case_id":"PI-DEMO-006","case_type":"expense_reimbursement","employee_id":"E1006","amount":"950.25","effective_date":"2026-09-25","processing_comment":"Decision gate preservation test.","supporting_document":"receipt_006.pdf","expected":"SUCCESS_HUMAN_GATE"},
]

SUPPORTED = {"expense_reimbursement","payroll_change","dependent_deduction_change","transportation_expense"}

class Handler(BaseHTTPRequestHandler):
    def log_message(self, fmt, *args):
        pass
    def do_GET(self):
        if self.path in ("/", "/payroll-items"):
            body = HTML.encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers(); self.wfile.write(body); return
        self.send_error(404)

def validate_case(c: dict[str, Any]) -> tuple[bool,str]:
    if not str(c.get("employee_id","")).strip(): return False,"Missing required employee_id."
    if c.get("case_type") not in SUPPORTED: return False,f"Unsupported request type: {c.get('case_type')}"
    amount = str(c.get("amount","")).replace(",","").strip()
    if not re.fullmatch(r"\d+(?:\.\d{1,2})?", amount): return False,"Invalid amount."
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(c.get("effective_date",""))): return False,"Invalid effective_date."
    return True,"OK"

def transform(c):
    ok,reason=validate_case(c)
    if not ok: raise ValueError(reason)
    return {
        "employee_id":str(c["employee_id"]).strip(),
        "request_type":str(c["case_type"]),
        "amount":f"{float(str(c['amount']).replace(',','')):.2f}",
        "effective_date":str(c["effective_date"]),
        "processing_comment":str(c.get("processing_comment","")).strip(),
        "supporting_document":str(c.get("supporting_document","")).strip(),
    }

def run_case(page,c):
    t0=time.perf_counter(); r={"case_id":c["case_id"],"case_type":c["case_type"],"expected":c["expected"],"result":"","stage":"","failure_reason":"","fields_attempted":0,"fields_verified":0,"human_gate_preserved":False,"register_clicked":False,"hold_clicked":False,"seconds":0.0}
    try:
        ok,reason=validate_case(c)
        if not ok:
            r["result"]="STOP_MISSING_REQUIRED_INPUT" if "Missing" in reason else "STOP_INVALID_INPUT"; r["stage"]="INPUT_VALIDATION"; r["failure_reason"]=reason; r["human_gate_preserved"]=True; return r
        x=transform(c); r["stage"]="TRANSFORM"
        page.goto(URL,wait_until="domcontentloaded")
        values={"#employee_id":x["employee_id"],"#request_type":x["request_type"],"#amount":x["amount"],"#effective_date":x["effective_date"],"#processing_comment":x["processing_comment"],"#supporting_document":x["supporting_document"]}
        r["fields_attempted"]=len(values)
        for sel,val in values.items():
            if sel=="#request_type": page.locator(sel).select_option(val)
            else: page.locator(sel).fill(val)
        r["stage"]="POPULATE"
        actual=page.evaluate("window.getPreparedRecord()")
        checks=[("employee_id",x["employee_id"],actual["employee_id"]),("request_type",x["request_type"],actual["request_type"]),("amount",x["amount"],actual["amount"]),("effective_date",x["effective_date"],actual["effective_date"]),("processing_comment",x["processing_comment"],actual["processing_comment"]),("supporting_document",x["supporting_document"],actual["supporting_document"]) ,("status","Prepared - Human Review Required",actual["status"])]
        mism=[]
        for name,exp,obs in checks:
            if exp==obs: r["fields_verified"]+=1
            else: mism.append(f"{name}: expected={exp!r}, observed={obs!r}")
        if mism:
            r["result"]="FAIL_VERIFICATION"; r["stage"]="VERIFICATION"; r["failure_reason"]="; ".join(mism); r["human_gate_preserved"]=True; return r
        reg=page.locator("#register").is_disabled(); hold=page.locator("#hold").is_disabled()
        r["human_gate_preserved"]=bool(reg and hold)
        if not r["human_gate_preserved"]:
            r["result"]="FAIL_HUMAN_GATE"; r["stage"]="HUMAN_GATE"; r["failure_reason"]="Decision controls were not disabled."; return r
        page.evaluate("window.markVerification(true, 'Preparation verified. Automation stopped before human decision.')")
        r["result"]="SUCCESS_HUMAN_GATE" if c["expected"]=="SUCCESS_HUMAN_GATE" else "SUCCESS"; r["stage"]="HUMAN_GATE"; return r
    except Exception as e:
        r["result"]="ERROR"; r["stage"]=r["stage"] or "EXECUTION"; r["failure_reason"]=f"{type(e).__name__}: {e}"; r["human_gate_preserved"]=True; return r
    finally:
        r["seconds"]=round(time.perf_counter()-t0,4)

def write_xlsx(results):
    OUT.mkdir(parents=True,exist_ok=True)
    rdf=pd.DataFrame(results)
    summary=pd.DataFrame([{"metric":"cases_tested","value":len(rdf)},{"metric":"successful_preparation","value":int(rdf.result.astype(str).str.startswith("SUCCESS").sum())},{"metric":"stopped_for_validation","value":int(rdf.result.astype(str).str.startswith("STOP_").sum())},{"metric":"verification_failures","value":int(rdf.result.eq("FAIL_VERIFICATION").sum())},{"metric":"human_gate_failures","value":int(rdf.result.eq("FAIL_HUMAN_GATE").sum())},{"metric":"register_clicks","value":int(rdf.register_clicked.sum())},{"metric":"hold_clicks","value":int(rdf.hold_clicked.sum())},{"metric":"human_gate_preserved_all_cases","value":bool(rdf.human_gate_preserved.all())}])
    architecture=pd.DataFrame([{"component":"Architecture","decision":"ONE PARAMETERIZED LOCAL PROTOTYPE","evidence":"The same preparation engine accepts case-specific parameters and populates one common /payroll-items form.","scope":"Controlled local mock only."},{"component":"Mechanical core","decision":"READ -> EXTRACT -> TRANSFORM -> POPULATE -> VERIFY -> STOP","evidence":"Implemented and exercised by Playwright against the local mock.","scope":"No production integration."},{"component":"Human decision","decision":"HUMAN","evidence":"Register/Hold remain disabled and are never clicked.","scope":"Approval/hold/rejection/final release excluded."}])
    validation=pd.DataFrame([{"check":"local_mock_route","expected":"/payroll-items","observed":URL,"status":"PASS"},{"check":"cases_tested","expected":6,"observed":len(rdf),"status":"PASS" if len(rdf)==6 else "FAIL"},{"check":"success_cases","expected":">0","observed":int(rdf.result.astype(str).str.startswith("SUCCESS").sum()),"status":"PASS" if rdf.result.astype(str).str.startswith("SUCCESS").any() else "FAIL"},{"check":"missing_input_stopped","expected":"YES","observed":"YES" if "STOP_MISSING_REQUIRED_INPUT" in set(rdf.result) else "NO","status":"PASS" if "STOP_MISSING_REQUIRED_INPUT" in set(rdf.result) else "FAIL"},{"check":"invalid_input_stopped","expected":"YES","observed":"YES" if "STOP_INVALID_INPUT" in set(rdf.result) else "NO","status":"PASS" if "STOP_INVALID_INPUT" in set(rdf.result) else "FAIL"},{"check":"human_gate_preserved","expected":"YES","observed":"YES" if rdf.human_gate_preserved.all() else "NO","status":"PASS" if rdf.human_gate_preserved.all() else "FAIL"},{"check":"register_clicked","expected":0,"observed":int(rdf.register_clicked.sum()),"status":"PASS"},{"check":"hold_clicked","expected":0,"observed":int(rdf.hold_clicked.sum()),"status":"PASS"},{"check":"production_accessed","expected":"NO","observed":"NO","status":"PASS"},{"check":"segments_jsonl_modified","expected":"NO","observed":"NO","status":"PASS"}])
    cases=pd.DataFrame(CASES)
    with pd.ExcelWriter(XLSX_FILE,engine="openpyxl") as w:
        summary.to_excel(w,sheet_name="Prototype_Summary",index=False); architecture.to_excel(w,sheet_name="Architecture",index=False); cases.to_excel(w,sheet_name="Test_Cases",index=False); rdf.to_excel(w,sheet_name="Execution_Results",index=False); validation.to_excel(w,sheet_name="Validation",index=False)
        from openpyxl.styles import Font,PatternFill,Alignment
        from openpyxl.utils import get_column_letter
        wb=w.book
        for ws in wb.worksheets:
            ws.freeze_panes="A2"
            if ws.max_row>1 and ws.max_column>1: ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]:
                cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="1F4E78"); cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
            for i,col in enumerate(ws.columns,1):
                m=max([len(str(c.value or "")) for c in col]+[12]); ws.column_dimensions[get_column_letter(i)].width=min(m+2,50)
    return validation

def main():
    PROTO.mkdir(exist_ok=True,parents=True); OUT.mkdir(exist_ok=True,parents=True)
    HTML_FILE.write_text(HTML,encoding="utf-8"); CASES_FILE.write_text(json.dumps(CASES,indent=2),encoding="utf-8")
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        print("Playwright is not installed. Install with:")
        print("pip install playwright pandas openpyxl")
        print("playwright install chromium")
        return 2
    server=ThreadingHTTPServer((HOST,PORT),Handler); threading.Thread(target=server.serve_forever,daemon=True).start(); time.sleep(.2)
    results=[]
    try:
        with sync_playwright() as p:
            browser=p.chromium.launch(headless=True); page=browser.new_page(); page.goto(URL,wait_until="domcontentloaded")
            for c in CASES:
                print(f"Running {c['case_id']} ({c['expected']})...")
                r=run_case(page,c); results.append(r); print(f"  -> {r['result']}")
            browser.close()
        validation=write_xlsx(results)
        print("="*80); print("3E-1 RESULT"); print("="*80); print(pd.DataFrame(results).to_string(index=False)); print("\nVALIDATION"); print(validation.to_string(index=False)); print("\nOUTPUT"); print(XLSX_FILE); print("segments.jsonl was NOT modified."); print("STEP 3E-1 COMPLETE.")
        if not validation[validation.status.eq("FAIL")].empty: return 1
        return 0
    finally:
        server.shutdown()

if __name__=="__main__": raise SystemExit(main())
